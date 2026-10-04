from __future__ import annotations

import json
import random
import re
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from faker import Faker
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


SEED = 20260929
OUTPUT_ROOT = Path(__file__).resolve().parent.parent / "generated"
DATA_DIR = OUTPUT_ROOT / "data"
PARQUET_DIR = DATA_DIR / "parquet"
DOCS_DIR = OUTPUT_ROOT / "documents"
SOP_DIR = DOCS_DIR / "Standard Operating Procedures"
MANUAL_DIR = DOCS_DIR / "Equipment Technical Manuals"

fake = Faker("en_US")
fake.seed_instance(SEED)
random.seed(SEED)

FACILITIES = [
    ("FAC-CIM", "Cimarron Operations Center", "OKC", "Oklahoma City", "OK", "Operations Center"),
    ("FAC-CRS", "Crossroads Energy Center", "OKC", "Oklahoma City", "OK", "Combined Cycle Plant"),
    ("FAC-ARB", "Arbuckle Service Center", "SOUTH", "Ardmore", "OK", "Service Center"),
    ("FAC-WAS", "Washita Peaking Station", "WEST", "Clinton", "OK", "Peaking Plant"),
    ("FAC-NCN", "North Canadian Substation", "OKC", "Yukon", "OK", "Transmission Substation"),
    ("FAC-SOO", "Sooner Operations Complex", "NORTH", "Stillwater", "OK", "Operations Complex"),
]

ASSET_CLASSES: dict[str, list[tuple[str, str, str]]] = {
    "Combined Cycle Plant": [
        ("GASTURB", "Gas Turbine", "Siemens Energy SGT6-5000F"),
        ("GENERATOR", "Generator", "GE 7A6"),
        ("TRANSFORMER", "Generator Step-Up Transformer", "Hitachi Energy GSU-230"),
        ("PUMP", "Boiler Feedwater Pump", "Flowserve WXH"),
        ("COMPRESSOR", "Instrument Air Compressor", "Atlas Copco GA90"),
        ("SWITCHGEAR", "Medium-Voltage Switchgear", "Eaton VacClad-W"),
    ],
    "Peaking Plant": [
        ("GASTURB", "Combustion Turbine", "GE LM6000"),
        ("GENERATOR", "Generator", "Brush BDAX"),
        ("TRANSFORMER", "Unit Auxiliary Transformer", "Virginia Transformer UAT-34"),
        ("FUELSKID", "Natural Gas Fuel Skid", "Emerson Fisher FS-600"),
        ("BATTERY", "Station Battery Bank", "EnerSys GU-45"),
    ],
    "Transmission Substation": [
        ("BREAKER", "230 kV Circuit Breaker", "Siemens 3AP1"),
        ("TRANSFORMER", "230/138 kV Autotransformer", "Hitachi Energy TXpert"),
        ("RELAY", "Protection Relay", "SEL-421"),
        ("BATTERY", "Station Battery Bank", "C&D LCR-25"),
        ("SWITCH", "Disconnect Switch", "Southern States EV-2"),
    ],
    "Operations Center": [
        ("UPS", "Critical Systems UPS", "Eaton 93PM"),
        ("HVAC", "Data Hall Air Handler", "Trane IntelliPak"),
        ("GENSET", "Emergency Generator", "Caterpillar C175-16"),
        ("FIRE", "Clean-Agent Fire Suppression Panel", "Kidde VS4-G"),
        ("SWITCHGEAR", "Facility Switchgear", "Square D Power-Zone 4"),
    ],
    "Service Center": [
        ("HVAC", "Rooftop Air Handler", "Trane Precedent"),
        ("GENSET", "Emergency Generator", "Cummins C500D6"),
        ("CRANE", "Bridge Crane", "Konecranes CXT"),
        ("COMPRESSOR", "Shop Air Compressor", "Ingersoll Rand R90"),
        ("CHARGER", "Fleet EV Charger", "ChargePoint Express Plus"),
    ],
    "Operations Complex": [
        ("UPS", "Operations UPS", "Eaton 93PM"),
        ("HVAC", "Operations Air Handler", "Carrier WeatherExpert"),
        ("GENSET", "Emergency Generator", "Caterpillar C18"),
        ("PUMP", "Fire Water Pump", "Aurora 912"),
        ("SWITCHGEAR", "Facility Switchgear", "Eaton Magnum DS"),
    ],
}

CREW_DEFS = [
    ("CREW-ELEC-01", "Central Electrical Maintenance", "Electrical"),
    ("CREW-MECH-01", "Rotating Equipment Maintenance", "Mechanical"),
    ("CREW-IC-01", "Instrumentation and Controls", "I&C"),
    ("CREW-SUB-01", "Substation Apparatus", "Substation"),
    ("CREW-FAC-01", "Facilities Services", "Facilities"),
    ("CREW-OPS-01", "Generation Operations", "Operations"),
    ("CREW-RELY-01", "Protection and Relay", "Relay"),
    ("CREW-CON-01", "Contractor Services", "Contractor"),
]

WORK_TEMPLATES = {
    "GASTURB": [
        ("Inspect combustion system and borescope accessible stages", "INSP", "MECHANICAL"),
        ("Investigate elevated bearing vibration alarm", "CM", "MECHANICAL"),
        ("Replace inlet filter differential pressure elements", "PM", "MECHANICAL"),
    ],
    "GENERATOR": [
        ("Perform generator winding insulation resistance test", "PM", "ELECTRICAL"),
        ("Investigate stator temperature deviation", "CM", "ELECTRICAL"),
        ("Inspect brush gear and grounding system", "INSP", "ELECTRICAL"),
    ],
    "TRANSFORMER": [
        ("Collect insulating oil sample for dissolved gas analysis", "PM", "ELECTRICAL"),
        ("Investigate main tank oil seepage", "CM", "ELECTRICAL"),
        ("Inspect bushings and infrared scan connections", "INSP", "ELECTRICAL"),
    ],
    "PUMP": [
        ("Inspect pump alignment, seals, and bearing condition", "PM", "MECHANICAL"),
        ("Correct mechanical seal leakage", "CM", "MECHANICAL"),
        ("Investigate low discharge pressure", "CM", "MECHANICAL"),
    ],
    "BREAKER": [
        ("Perform breaker timing and contact resistance tests", "PM", "ELECTRICAL"),
        ("Investigate mechanism low-pressure alarm", "CM", "ELECTRICAL"),
        ("Inspect interrupters, bushings, and operating cabinet", "INSP", "ELECTRICAL"),
    ],
    "RELAY": [
        ("Perform protection relay calibration and trip verification", "PM", "I&C"),
        ("Review relay target following system operation", "CM", "I&C"),
        ("Validate settings against approved coordination study", "INSP", "I&C"),
    ],
}

DEFAULT_TEMPLATES = [
    ("Complete quarterly condition inspection", "PM", "FACILITIES"),
    ("Investigate abnormal operating indication", "CM", "FACILITIES"),
    ("Perform annual functional test", "PM", "ELECTRICAL"),
    ("Correct deficiency identified during operator rounds", "CM", "MECHANICAL"),
]

PUMP_LOW_PRESSURE_CASES = [
    {
        "description": "Pump discharge pressure below normal operating range",
        "details": "Operators reported reduced header pressure with stable motor current. Inspection found a partially obstructed suction strainer; the strainer was cleaned and normal pressure was restored.",
        "problem_code": "DEGRADED_OUTPUT",
        "cause_code": "CONTAMINATION",
        "remedy_code": "CLEANED",
    },
    {
        "description": "Investigate reduced flow from process water pump",
        "details": "The pump could not sustain expected process flow during peak demand. Testing identified excessive impeller wear; the rotating assembly was repaired and performance was verified against the pump curve.",
        "problem_code": "DEGRADED_OUTPUT",
        "cause_code": "WEAR",
        "remedy_code": "REPAIRED",
    },
    {
        "description": "Centrifugal pump develops insufficient discharge head",
        "details": "Discharge head declined while suction pressure remained available. Coupling checks found shaft misalignment after recent foundation work; alignment was corrected and the unit was retested.",
        "problem_code": "DEGRADED_OUTPUT",
        "cause_code": "MISALIGNMENT",
        "remedy_code": "ADJUSTED",
    },
    {
        "description": "Intermittent low-pressure condition on transfer pump",
        "details": "Pressure fluctuated during starts and operators heard gravel-like noise near the casing. Air was entering through a degraded suction-side seal; the seal was replaced and the pump was vented before testing.",
        "problem_code": "DEGRADED_OUTPUT",
        "cause_code": "WEAR",
        "remedy_code": "REPLACED",
    },
    {
        "description": "Pump output has fallen below recent baseline",
        "details": "Trend review showed a gradual loss of developed pressure without an electrical anomaly. Deposits were found in the impeller passages; the casing and impeller were cleaned and output returned to baseline.",
        "problem_code": "DEGRADED_OUTPUT",
        "cause_code": "CONTAMINATION",
        "remedy_code": "CLEANED",
    },
    {
        "description": "Troubleshoot weak discharge from cooling-water pump",
        "details": "The cooling loop received inadequate flow even though the pump remained in service. A worn wear ring increased internal recirculation; the wear ring was replaced and post-maintenance readings met acceptance criteria.",
        "problem_code": "DEGRADED_OUTPUT",
        "cause_code": "WEAR",
        "remedy_code": "REPLACED",
    },
]

PUMP_OPERATING_CONTEXTS = [
    "startup",
    "peak process demand",
    "operator rounds",
    "weekly functional testing",
    "post-maintenance monitoring",
    "steady-state operation",
]

DOCUMENTS = [
    {
        "docnum": "SOP-ELEC-001",
        "library": "Standard Operating Procedures",
        "title": "Electrical Lockout and Tagout for Facility Equipment",
        "revision": "4",
        "owner": "Electrical Safety",
        "asset_class": "ALL",
        "purpose": "Establish the minimum steps for isolating hazardous electrical energy before inspection, maintenance, or testing.",
        "sections": [
            ("Scope", ["Applies to qualified employees and contractors working on facility electrical equipment at 50 volts or greater.", "Use the approved switching order when utility-system switching is involved."]),
            ("Required PPE and tools", ["Arc-rated clothing selected from the current arc-flash label", "Voltage-rated gloves with leather protectors", "Approved voltage detector and personal locks", "One-line diagram and equipment-specific isolation checklist"]),
            ("Procedure", ["Review the work order, boundaries, and one-line diagram during the pre-job brief.", "Identify every normal, alternate, stored, and backfeed energy source.", "Open isolation devices, apply personal locks and tags, and release stored energy.", "Verify the test instrument on a known live source, test phase-to-phase and phase-to-ground, then reverify the instrument.", "Install protective grounds when required by the switching order.", "Document isolation points and independent verification in the work order."]),
            ("Restoration", ["Inspect the work area and account for tools, grounds, and personnel.", "Each employee removes their personal lock; exceptions require the approved abandoned-lock process.", "Notify affected personnel and restore equipment using the approved operating sequence."]),
        ],
    },
    {
        "docnum": "SOP-MECH-004",
        "library": "Standard Operating Procedures",
        "title": "Rotating Equipment Condition Inspection",
        "revision": "3",
        "owner": "Maintenance Engineering",
        "asset_class": "PUMP",
        "purpose": "Standardize condition inspections for pumps, compressors, motors, and coupled rotating equipment.",
        "sections": [
            ("Prerequisites", ["Review baseline vibration and temperature data.", "Confirm guards are installed and the equipment is in a stable operating condition."]),
            ("Inspection route", ["Record suction and discharge pressure.", "Measure horizontal, vertical, and axial vibration at each bearing.", "Record bearing and seal temperatures.", "Inspect lubrication level, coupling guard, baseplate, anchor bolts, and seal leakage.", "Listen for cavitation, rubbing, looseness, or abnormal gear noise."]),
            ("Acceptance criteria", ["Escalate vibration increases greater than 25 percent from baseline.", "Create a corrective work order for active oil leakage, loose fasteners, or degraded guards.", "Notify Operations immediately for rapid temperature rise or suspected bearing damage."]),
        ],
    },
    {
        "docnum": "SOP-SUB-007",
        "library": "Standard Operating Procedures",
        "title": "Substation Breaker Inspection and Diagnostic Testing",
        "revision": "2",
        "owner": "Substation Maintenance",
        "asset_class": "BREAKER",
        "purpose": "Define inspection and diagnostic testing requirements for high-voltage circuit breakers.",
        "sections": [
            ("Safety controls", ["Obtain an approved clearance and establish the minimum approach boundary.", "Ground isolated terminals according to the station switching order.", "Disable remote and automatic operation before mechanism work."]),
            ("Inspection", ["Inspect bushings, interrupters, cabinets, heaters, wiring, and grounding.", "Record gas density or hydraulic pressure and ambient temperature.", "Check mechanism lubrication and evidence of water or animal intrusion."]),
            ("Diagnostic tests", ["Measure contact resistance on each pole.", "Perform close, open, and close-open timing tests.", "Verify trip and close coil current signatures.", "Record results in the work order and compare with manufacturer limits and previous tests."]),
        ],
    },
    {
        "docnum": "SOP-OPS-011",
        "library": "Standard Operating Procedures",
        "title": "Abnormal Equipment Condition Response",
        "revision": "5",
        "owner": "System Operations",
        "asset_class": "ALL",
        "purpose": "Provide a consistent response to alarms, leaks, unusual noise, odor, heat, vibration, or degraded equipment condition.",
        "sections": [
            ("Immediate actions", ["Protect personnel and establish a safe boundary.", "Notify the control room or responsible operator.", "Place equipment in a safe state when directed and within operating authority.", "Capture alarm, trend, photograph, and operating-condition evidence."]),
            ("Work management", ["Create a priority work order using the approved priority matrix.", "Reference the parent incident or operator log entry.", "Identify temporary controls and required follow-up date.", "Escalate conditions with safety, environmental, or reliability consequences."]),
            ("Turnover", ["Brief the receiving shift on equipment status and temporary controls.", "Confirm ownership for engineering review and corrective maintenance."]),
        ],
    },
    {
        "docnum": "SOP-ENV-014",
        "library": "Standard Operating Procedures",
        "title": "Oil Spill Prevention and Initial Response",
        "revision": "3",
        "owner": "Environmental Compliance",
        "asset_class": "TRANSFORMER",
        "purpose": "Define initial response actions for transformer oil, lubricating oil, and fuel releases.",
        "sections": [
            ("Response", ["Stop work and protect personnel.", "Stop the source when safe to do so.", "Protect drains, soil, and surface water with available spill-control materials.", "Notify Operations and Environmental Compliance.", "Estimate quantity, affected media, and containment status without delaying control actions."]),
            ("Documentation", ["Record time discovered, weather, source equipment, photographs, and materials used.", "Create follow-up work orders for cleanup, repair, and inspection.", "Do not make external notifications unless assigned by Environmental Compliance."]),
        ],
    },
    {
        "docnum": "SOP-CMMS-021",
        "library": "Standard Operating Procedures",
        "title": "Work Order Planning, Execution, and Closeout",
        "revision": "6",
        "owner": "Work Management",
        "asset_class": "ALL",
        "purpose": "Define quality requirements for planning, scheduling, execution, and closeout of maintenance work orders.",
        "sections": [
            ("Planning", ["Confirm asset, location, failure hierarchy, and priority.", "Define labor crafts, estimated hours, materials, permits, procedures, and operational impacts.", "Attach drawings and applicable technical documents."]),
            ("Execution", ["Record actual start and finish times, labor, materials, measurements, and as-found/as-left condition.", "Document scope changes and obtain approval before performing emergent additional work."]),
            ("Closeout", ["Enter failure class, problem, cause, and remedy for corrective work.", "Record test results and follow-up recommendations.", "Confirm the asset is returned to an approved status and route the record for supervisor review."]),
        ],
    },
    {
        "docnum": "MAN-GT-100",
        "library": "Equipment Technical Manuals",
        "title": "SGT6-Class Gas Turbine Operations and Routine Maintenance Guide",
        "revision": "1",
        "owner": "Generation Engineering",
        "asset_class": "GASTURB",
        "purpose": "Provide demo operating limits, inspection points, and routine maintenance guidance for an industrial gas turbine package.",
        "sections": [
            ("System overview", ["The package includes inlet filtration, compressor, combustion system, turbine, accessory gearbox, enclosure ventilation, fire protection, and controls.", "The information is synthetic and must not replace manufacturer instructions."]),
            ("Operator checks", ["Monitor bearing vibration, exhaust temperature spread, lube oil pressure, inlet differential pressure, and enclosure ventilation.", "Investigate trends rather than relying only on alarm thresholds."]),
            ("Routine maintenance", ["Inspect inlet filters and drain systems.", "Check fire-detection devices and enclosure dampers.", "Sample lubricating oil and inspect magnetic chip detectors.", "Borescope compressor and turbine sections at the approved interval."]),
            ("Troubleshooting", ["High exhaust spread: validate instrumentation, fuel distribution, and combustion hardware.", "Increasing vibration: verify sensor integrity, operating condition, alignment history, and bearing condition.", "Low lube pressure: confirm oil level, temperature, filter differential pressure, and pump operation."]),
        ],
    },
    {
        "docnum": "MAN-XFMR-210",
        "library": "Equipment Technical Manuals",
        "title": "Power Transformer Inspection and Maintenance Guide",
        "revision": "2",
        "owner": "Substation Engineering",
        "asset_class": "TRANSFORMER",
        "purpose": "Provide routine inspection, oil testing, and maintenance guidance for liquid-filled power transformers.",
        "sections": [
            ("Inspection", ["Check oil level, pressure, temperature, cooling-system status, bushings, arresters, grounding, and evidence of leakage.", "Use infrared inspection under representative load."]),
            ("Oil program", ["Collect samples using clean, dry equipment and purge the sample valve.", "Trend dissolved gases, moisture, dielectric strength, acidity, and inhibitor content.", "Escalate significant rate-of-change even when absolute values remain below alarm criteria."]),
            ("Maintenance", ["Test alarms, fans, pumps, and control power.", "Inspect tap-changer drive and counter operation.", "Correct active leaks using an engineering-approved repair method."]),
        ],
    },
    {
        "docnum": "MAN-CB-230",
        "library": "Equipment Technical Manuals",
        "title": "230 kV Circuit Breaker Mechanism and Diagnostic Test Guide",
        "revision": "1",
        "owner": "Substation Engineering",
        "asset_class": "BREAKER",
        "purpose": "Provide synthetic service guidance for stored-energy high-voltage circuit breakers.",
        "sections": [
            ("Mechanism", ["The spring-hydraulic mechanism stores energy for close and trip operations.", "Interlocks prevent operation outside acceptable pressure and energy conditions."]),
            ("Service checks", ["Inspect heaters, seals, wiring, linkages, lubrication, and operation counters.", "Verify local/remote selector and anti-pump function.", "Record minimum operating voltage and coil current when directed by the job plan."]),
            ("Limits", ["Compare timing and resistance by pole and against previous results.", "Remove the breaker from service for evidence of internal distress, unstable gas density, or mechanism binding."]),
        ],
    },
    {
        "docnum": "MAN-UPS-310",
        "library": "Equipment Technical Manuals",
        "title": "Critical Systems UPS Operations and Battery Maintenance Guide",
        "revision": "3",
        "owner": "Facilities Engineering",
        "asset_class": "UPS",
        "purpose": "Describe normal operation, bypass controls, alarm response, and battery maintenance for critical facility UPS systems.",
        "sections": [
            ("Operating modes", ["Normal mode supplies the load through the rectifier and inverter.", "Static bypass supports automatic transfer during overload or inverter fault.", "Maintenance bypass isolates power electronics and requires an approved switching procedure."]),
            ("Monthly checks", ["Review active and historical alarms.", "Record input, output, bypass, load, and battery values.", "Inspect ventilation and room temperature.", "Check for battery swelling, corrosion, or leakage without contacting energized parts."]),
            ("Alarm response", ["Do not transfer modes solely to clear an alarm.", "Assess redundancy and load before maintenance.", "Escalate battery ground, overtemperature, fan failure, and bypass-unavailable conditions."]),
        ],
    },
    {
        "docnum": "MAN-PUMP-410",
        "library": "Equipment Technical Manuals",
        "title": "Horizontal Centrifugal Pump Maintenance Guide",
        "revision": "2",
        "owner": "Maintenance Engineering",
        "asset_class": "PUMP",
        "purpose": "Provide installation checks, lubrication practices, seal inspection, alignment, and troubleshooting guidance.",
        "sections": [
            ("Installation checks", ["Verify foundation, grout, piping support, shaft rotation, and soft foot.", "Perform final alignment after piping connection and at operating temperature when required."]),
            ("Maintenance", ["Use the specified lubricant and avoid overgreasing.", "Trend vibration, bearing temperature, seal leakage, suction pressure, and discharge pressure.", "Inspect coupling elements and guards at planned intervals."]),
            ("Troubleshooting", ["Cavitation: verify available suction head, strainers, valve position, and process temperature.", "Seal leakage: inspect alignment, runout, flush plan, and seal faces.", "High vibration: check imbalance, looseness, misalignment, resonance, and bearing condition."]),
        ],
    },
    {
        "docnum": "MAN-GEN-510",
        "library": "Equipment Technical Manuals",
        "title": "Emergency Generator Inspection and Load Test Guide",
        "revision": "4",
        "owner": "Facilities Engineering",
        "asset_class": "GENSET",
        "purpose": "Define routine inspection and operational testing for standby diesel generator systems.",
        "sections": [
            ("Pre-start inspection", ["Check fuel, coolant, oil, starting batteries, charger, belts, hoses, exhaust, enclosure, and leaks.", "Verify generator breaker position and test prerequisites."]),
            ("Operational test", ["Start from the approved local or remote test control.", "Record voltage, frequency, current, power, oil pressure, coolant temperature, and battery charging.", "Apply facility load or load bank according to the approved test plan."]),
            ("Closeout", ["Allow unloaded cooldown when required.", "Return controls and breakers to their normal standby position.", "Record run hours, fuel level, deficiencies, and follow-up work orders."]),
        ],
    },
    {
        "docnum": "SOP-SAFE-025",
        "library": "Standard Operating Procedures",
        "title": "Permit-Required Confined Space Entry",
        "revision": "3",
        "owner": "Safety and Health",
        "asset_class": "ALL",
        "purpose": "Define controls for entry into permit-required confined spaces at generation, substation, and service-center facilities.",
        "sections": [
            ("Entry planning", ["Identify hazards, isolation points, authorized entrants, attendant, entry supervisor, and rescue method.", "Review atmospheric hazards, engulfment, configuration, stored energy, and work introduced by the entry."]),
            ("Entry controls", ["Isolate energy and material sources using approved lockout and blanking methods.", "Test oxygen, flammable gases, and potential toxic contaminants before and continuously during entry.", "Maintain ventilation, communications, entrant accountability, and retrieval equipment as specified by the permit."]),
            ("Termination", ["Evacuate immediately for alarms, symptoms, loss of ventilation, communication failure, or changing conditions.", "Close the permit, document readings and deviations, and restore the space only after the entry supervisor confirms completion."]),
        ],
    },
    {
        "docnum": "SOP-ELEC-029",
        "library": "Standard Operating Procedures",
        "title": "Medium-Voltage Switching and Clearance Control",
        "revision": "4",
        "owner": "System Operations",
        "asset_class": "SWITCHGEAR",
        "purpose": "Establish preparation, independent verification, execution, and closeout requirements for medium-voltage switching.",
        "sections": [
            ("Preparation", ["Use an approved switching order tied to the current one-line diagram.", "Confirm equipment identification, operating authority, system conditions, protective-device status, and contingency plan.", "Conduct a place-keeping review and independent verification before execution."]),
            ("Execution", ["Communicate each step using three-part repeat-back.", "Perform steps in sequence and mark each step complete immediately.", "Stop for unexpected indication, abnormal sound, failed operation, or mismatch between field labeling and the switching order."]),
            ("Clearance and restoration", ["Establish visible isolation, verify absence of voltage, and apply grounds where required.", "Account for personnel, tools, grounds, and tags before release.", "Document final equipment position and notify affected operating authorities."]),
        ],
    },
    {
        "docnum": "SOP-SAFE-033",
        "library": "Standard Operating Procedures",
        "title": "Hot Work Permit and Fire Watch",
        "revision": "2",
        "owner": "Fire Protection and Safety",
        "asset_class": "ALL",
        "purpose": "Define fire-prevention controls for welding, cutting, brazing, grinding, and other spark- or heat-producing work.",
        "sections": [
            ("Permit review", ["Evaluate combustible construction, process hazards, drains, wall penetrations, adjacent levels, and fire-protection impairments.", "Move combustibles or protect them with listed blankets and barriers.", "Verify suitable extinguishers and establish the fire-watch duration."]),
            ("Work controls", ["Inspect leads, hoses, cylinders, regulators, grounding, and ventilation.", "Maintain a trained fire watch with no conflicting duties.", "Stop work for gas alarms, changing conditions, loss of fire protection, or permit expiration."]),
            ("Post-work inspection", ["Inspect the work area and adjacent concealed spaces at the required intervals.", "Close the permit only after the final fire-watch inspection is complete.", "Record impairments, alarms, and follow-up corrective work."]),
        ],
    },
    {
        "docnum": "SOP-DC-037",
        "library": "Standard Operating Procedures",
        "title": "Station Battery Inspection and Capacity Trending",
        "revision": "3",
        "owner": "Protection and Control",
        "asset_class": "BATTERY",
        "purpose": "Standardize inspection, measurement, and trending of station battery systems supporting protection and control loads.",
        "sections": [
            ("Safety", ["Review battery chemistry, shock and arc hazards, ventilation, eyewash access, and spill-response provisions.", "Remove conductive jewelry and use insulated tools and specified PPE."]),
            ("Inspection", ["Record charger output, bank voltage, cell or unit voltage, temperature, connection condition, electrolyte level where applicable, and active alarms.", "Inspect racks, seismic restraints, ventilation, grounding, corrosion, swelling, cracking, and leakage."]),
            ("Trending", ["Compare readings with baseline and previous inspections.", "Escalate abnormal voltage spread, temperature rise, corrosion, impedance increase, or repeated charger alarms.", "Create corrective work orders with affected cell identifiers and measured values."]),
        ],
    },
    {
        "docnum": "SOP-FAC-041",
        "library": "Standard Operating Procedures",
        "title": "Critical HVAC Alarm Response and Temporary Cooling",
        "revision": "2",
        "owner": "Facilities Operations",
        "asset_class": "HVAC",
        "purpose": "Define response priorities for HVAC failures affecting control rooms, relay rooms, data halls, and battery rooms.",
        "sections": [
            ("Assessment", ["Confirm alarm validity, room temperature, humidity, equipment status, redundant-unit availability, and affected critical loads.", "Notify the responsible operator and facilities supervisor."]),
            ("Stabilization", ["Start standby equipment using the approved operating sequence.", "Deploy temporary cooling only after electrical load, condensate routing, airflow, and fire egress are reviewed.", "Increase temperature monitoring frequency and establish escalation thresholds."]),
            ("Recovery", ["Restore automatic controls and redundancy after repair.", "Review trend data for thermal impact to batteries, relays, servers, and controls.", "Document temporary equipment removal and required follow-up."]),
        ],
    },
    {
        "docnum": "SOP-RIG-045",
        "library": "Standard Operating Procedures",
        "title": "Critical Lift Planning and Rigging Inspection",
        "revision": "1",
        "owner": "Maintenance and Construction",
        "asset_class": "CRANE",
        "purpose": "Define planning and execution controls for critical lifts involving utility equipment, occupied areas, or limited clearances.",
        "sections": [
            ("Lift classification", ["Classify the lift using load weight, capacity percentage, tandem-crane use, energized-equipment proximity, consequence, and load-path criteria.", "Assign a qualified lift director, operator, rigger, and signal person."]),
            ("Lift plan", ["Verify load weight, center of gravity, crane configuration, ground conditions, rigging capacity, clearances, weather limits, and exclusion zone.", "Inspect slings, hardware, lifting points, hooks, brakes, indicators, and required certifications."]),
            ("Execution", ["Conduct a pre-lift briefing and test lift.", "Use one designated signal person unless an emergency stop is given.", "Stop for loss of communication, unexpected load movement, changing weather, or unauthorized entry into the exclusion zone."]),
        ],
    },
    {
        "docnum": "MAN-AIR-610",
        "library": "Equipment Technical Manuals",
        "title": "Oil-Injected Rotary Screw Air Compressor Maintenance Guide",
        "revision": "2",
        "owner": "Maintenance Engineering",
        "asset_class": "COMPRESSOR",
        "purpose": "Provide routine inspection, service, and troubleshooting guidance for plant and shop rotary screw air compressors.",
        "sections": [
            ("Operating checks", ["Record discharge pressure, oil temperature, separator differential pressure, motor current, loaded hours, and unloaded hours.", "Inspect condensate drains, coolers, ventilation, belts or coupling, hoses, and evidence of leakage."]),
            ("Scheduled service", ["Replace inlet filters, oil filters, separators, and lubricant at approved intervals.", "Clean coolers using methods that do not damage fins.", "Verify safety shutdowns and pressure-relief protection according to the job plan."]),
            ("Troubleshooting", ["High temperature: check oil level, cooler cleanliness, ventilation, thermostatic valve, and ambient conditions.", "Low capacity: inspect inlet restriction, leaks, control valves, and system demand.", "Excessive oil carryover: evaluate separator condition, scavenge line, oil level, and operating pressure."]),
        ],
    },
    {
        "docnum": "MAN-SWG-620",
        "library": "Equipment Technical Manuals",
        "title": "Medium-Voltage Metal-Clad Switchgear Maintenance Guide",
        "revision": "3",
        "owner": "Electrical Engineering",
        "asset_class": "SWITCHGEAR",
        "purpose": "Provide inspection, cleaning, mechanical service, and electrical testing guidance for metal-clad switchgear.",
        "sections": [
            ("Inspection", ["Inspect insulation, shutters, racking mechanism, grounding contacts, heaters, wiring, barriers, and evidence of tracking or contamination.", "Verify labels, mimic bus, interlocks, and compartment hardware."]),
            ("Maintenance", ["Clean using manufacturer-approved methods and materials.", "Lubricate only designated points with the specified lubricant.", "Exercise racking and interlocks with equipment de-energized and in the approved test configuration."]),
            ("Testing", ["Perform insulation resistance, contact resistance, breaker timing, and protective-device functional tests as specified by the job plan.", "Compare phase values and historical trends before returning equipment to service."]),
        ],
    },
    {
        "docnum": "MAN-RELAY-630",
        "library": "Equipment Technical Manuals",
        "title": "Digital Protection Relay Testing and Configuration Control Guide",
        "revision": "2",
        "owner": "Protection Engineering",
        "asset_class": "RELAY",
        "purpose": "Define configuration control, functional testing, and records requirements for digital protection relays.",
        "sections": [
            ("Configuration control", ["Obtain the approved settings file and checksum from the controlled repository.", "Record firmware, serial number, active settings group, time source, and communication status before changes.", "Back up the as-found configuration."]),
            ("Testing", ["Isolate trip outputs using the approved test plan.", "Inject current and voltage to verify pickup, timing, logic, metering, alarms, and output contacts.", "Perform an end-to-end or trip-path test when required by the maintenance basis."]),
            ("Closeout", ["Load only the approved configuration and independently verify checksum and active settings group.", "Restore trip circuits, communications, time synchronization, and alarms.", "Attach test results and as-left configuration to the work order."]),
        ],
    },
    {
        "docnum": "MAN-CRANE-640",
        "library": "Equipment Technical Manuals",
        "title": "Overhead Bridge Crane Inspection and Maintenance Guide",
        "revision": "1",
        "owner": "Facilities Engineering",
        "asset_class": "CRANE",
        "purpose": "Provide pre-use, periodic, and planned maintenance guidance for powered overhead bridge cranes.",
        "sections": [
            ("Pre-use inspection", ["Check controls, upper limit, brakes, hook, latch, wire rope or chain, warning devices, and travel path.", "Remove the crane from service for damaged load-bearing components or failed safety devices."]),
            ("Periodic inspection", ["Inspect runway, end trucks, wheels, gears, couplings, drums, sheaves, collectors, festoon systems, and structural connections.", "Measure hook throat opening and twist against baseline values."]),
            ("Maintenance", ["Apply lockout controls before entering the bridge or working on moving components.", "Use specified lubricants and adjustment limits.", "Perform an operational test after maintenance and a rated-load test when required by the repair scope."]),
        ],
    },
    {
        "docnum": "MAN-FIRE-650",
        "library": "Equipment Technical Manuals",
        "title": "Electric and Diesel Fire Pump Inspection and Test Guide",
        "revision": "3",
        "owner": "Fire Protection Engineering",
        "asset_class": "PUMP",
        "purpose": "Provide routine inspection and operational test guidance for facility fire-pump systems.",
        "sections": [
            ("Inspection", ["Verify controller status, power or fuel supply, suction and discharge pressure, valve position, room heat, and absence of leakage.", "Inspect batteries, chargers, cooling systems, exhaust, and engine fluids on diesel units."]),
            ("No-flow test", ["Start automatically using the approved pressure-drop method.", "Record churn pressure, suction pressure, start pressure, run time, and abnormal conditions.", "Do not leave the controller in manual or off after testing."]),
            ("Flow test", ["Use calibrated instruments and approved discharge arrangements.", "Record flow, net pressure, speed, voltage, current, and engine parameters at required points.", "Compare results with the acceptance curve and previous tests."]),
        ],
    },
    {
        "docnum": "MAN-EVC-660",
        "library": "Equipment Technical Manuals",
        "title": "Fleet Electric Vehicle Charging System Service Guide",
        "revision": "1",
        "owner": "Fleet Electrification",
        "asset_class": "CHARGER",
        "purpose": "Provide inspection, isolation, firmware, connector, and functional-test guidance for fleet charging equipment.",
        "sections": [
            ("Inspection", ["Inspect enclosure, cable management, connector pins, strain relief, impact damage, seals, labels, bollards, and ventilation.", "Review charger, network, and ground-fault event logs."]),
            ("Service", ["Isolate AC and DC energy sources using approved procedures.", "Verify absence of voltage and observe capacitor discharge time.", "Use only approved firmware and configuration packages with documented versions."]),
            ("Functional test", ["Verify authentication, communication, emergency stop, ground-fault protection, connector lock, charging initiation, and controlled stop.", "Record delivered voltage, current, power, and session result.", "Return cable management and protective barriers to service."]),
        ],
    },
]


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def iso(value: datetime | date | None) -> str | None:
    if value is None:
        return None
    return value.isoformat()


def write_table(name: str, rows: list[dict[str, Any]]) -> None:
    frame = pd.DataFrame(rows)
    frame.to_parquet(PARQUET_DIR / f"{name}.parquet", index=False)


def generate_data() -> dict[str, list[dict[str, Any]]]:
    now = datetime(2026, 9, 29, 12, 0, tzinfo=UTC)
    tables: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for siteid, name, region, city, state, facility_type in FACILITIES:
        tables["sites"].append(
            {
                "siteid": siteid,
                "orgid": "OGE",
                "description": name,
                "facility_type": facility_type,
                "region": region,
                "city": city,
                "state": state,
                "timezone": "America/Chicago",
                "active": True,
                "data_classification": "SYNTHETIC",
            }
        )
        for index, area in enumerate(("Administration", "Electrical", "Mechanical", "Operations", "Warehouse"), 1):
            tables["locations"].append(
                {
                    "location": f"{siteid}-{index:02d}",
                    "siteid": siteid,
                    "orgid": "OGE",
                    "description": f"{name} - {area}",
                    "type": "OPERATING",
                    "status": "OPERATING",
                    "parent": siteid if index > 1 else None,
                }
            )

    for crewnum, description, craft in CREW_DEFS:
        tables["crews"].append(
            {
                "crewnum": crewnum,
                "description": description,
                "default_craft": craft,
                "shift": random.choice(["DAY", "DAY", "ROTATING"]),
                "active": True,
            }
        )

    crafts = ["ELECTRICAL", "MECHANICAL", "I&C", "OPERATIONS", "FACILITIES", "RELAY", "PLANNER"]
    for labor_index in range(1, 65):
        first, last = fake.first_name(), fake.last_name()
        craft = random.choice(crafts)
        eligible_crews = [c[0] for c in CREW_DEFS if craft.split()[0].lower() in (c[2].lower(), c[1].lower())]
        tables["labor"].append(
            {
                "laborcode": f"L{labor_index:04d}",
                "displayname": f"{first} {last}",
                "craft": craft,
                "skilllevel": random.choice(["APPRENTICE", "JOURNEY", "SENIOR", "LEAD"]),
                "crewnum": random.choice(eligible_crews or [c[0] for c in CREW_DEFS]),
                "home_siteid": random.choice(FACILITIES)[0],
                "shift": random.choice(["DAY", "DAY", "ROTATING"]),
                "active": True,
            }
        )

    asset_counter = 1
    for facility in FACILITIES:
        siteid, site_name, _, _, _, facility_type = facility
        classes = ASSET_CLASSES[facility_type]
        count = 48 if facility_type in ("Combined Cycle Plant", "Transmission Substation") else 32
        locations = [row["location"] for row in tables["locations"] if row["siteid"] == siteid]
        for _ in range(count):
            asset_class, description, model = random.choice(classes)
            commissioned = date(2001, 1, 1) + timedelta(days=random.randint(0, 8500))
            assetnum = f"A{asset_counter:06d}"
            tables["assets"].append(
                {
                    "assetnum": assetnum,
                    "siteid": siteid,
                    "orgid": "OGE",
                    "location": random.choice(locations),
                    "description": f"{description} {asset_counter % 20 + 1:02d}",
                    "asset_class": asset_class,
                    "manufacturer": model.split()[0],
                    "model": model,
                    "serialnum": f"{asset_class[:3]}-{commissioned.year}-{random.randint(10000, 99999)}",
                    "status": random.choices(["OPERATING", "OPERATING", "OPERATING", "STANDBY"], [70, 10, 10, 10])[0],
                    "criticality": random.choices(["1-CRITICAL", "2-HIGH", "3-MEDIUM", "4-LOW"], [15, 30, 40, 15])[0],
                    "commission_date": iso(commissioned),
                    "replacement_value_usd": round(random.uniform(20_000, 7_500_000), 2),
                    "parent": None,
                    "synthetic_source": site_name,
                }
            )
            asset_counter += 1

    inventory_names = [
        ("BRG", "Bearing assembly", "EA", 180.0, 2200.0),
        ("FILT", "Filter element", "EA", 22.0, 480.0),
        ("GASK", "Industrial gasket kit", "KIT", 35.0, 620.0),
        ("LUBE", "Equipment lubricant", "GAL", 18.0, 160.0),
        ("FUSE", "Control fuse", "EA", 8.0, 95.0),
        ("RELAY", "Auxiliary control relay", "EA", 45.0, 390.0),
        ("SEAL", "Mechanical seal kit", "KIT", 240.0, 3800.0),
        ("CONT", "Power contactor", "EA", 160.0, 2100.0),
        ("SENS", "Temperature sensor", "EA", 65.0, 760.0),
        ("BOLT", "High-strength fastener kit", "KIT", 18.0, 210.0),
    ]
    for item_index in range(1, 151):
        prefix, description, unit, low, high = random.choice(inventory_names)
        tables["inventory_items"].append(
            {
                "itemnum": f"{prefix}-{item_index:05d}",
                "description": f"{description}, specification {item_index:03d}",
                "commodity_group": prefix,
                "orderunit": unit,
                "issueunit": unit,
                "unitcost_usd": round(random.uniform(low, high), 2),
                "manufacturer": random.choice(["Eaton", "SKF", "Parker", "Emerson", "ABB", "Siemens", "Timken"]),
                "catalogcode": f"CAT-{random.randint(100000, 999999)}",
                "condition_enabled": prefix in {"BRG", "RELAY", "CONT", "SENS"},
            }
        )

    job_plans: list[dict[str, Any]] = []
    class_set = sorted({row["asset_class"] for row in tables["assets"]})
    for jp_index, asset_class in enumerate(class_set, 1):
        templates = WORK_TEMPLATES.get(asset_class, DEFAULT_TEMPLATES)
        for template_index, (description, worktype, craft) in enumerate(templates[:2], 1):
            job_plans.append(
                {
                    "jpnum": f"JP-{asset_class[:5]}-{template_index:02d}",
                    "description": description,
                    "asset_class": asset_class,
                    "worktype": worktype,
                    "status": "ACTIVE",
                    "revision": random.randint(1, 5),
                    "estimated_hours": round(random.uniform(2, 24), 1),
                    "crew_size": random.randint(1, 4),
                    "lead_craft": craft,
                    "procedure_docnum": next((d["docnum"] for d in DOCUMENTS if d["asset_class"] in (asset_class, "ALL")), "SOP-CMMS-021"),
                }
            )
    tables["job_plans"] = job_plans

    statuses = ["WAPPR", "APPR", "WMATL", "SCHED", "INPRG", "COMP", "CLOSE", "CAN"]
    completed_statuses = {"COMP", "CLOSE"}
    asset_lookup = {row["assetnum"]: row for row in tables["assets"]}
    labor_by_craft: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for labor in tables["labor"]:
        labor_by_craft[labor["craft"]].append(labor)

    for wo_index in range(1, 1501):
        asset = random.choice(tables["assets"])
        templates = WORK_TEMPLATES.get(asset["asset_class"], DEFAULT_TEMPLATES)
        description, default_type, lead_craft = random.choice(templates)
        pump_low_pressure_case = (
            PUMP_LOW_PRESSURE_CASES[wo_index % len(PUMP_LOW_PRESSURE_CASES)]
            if description == "Investigate low discharge pressure"
            else None
        )
        if pump_low_pressure_case:
            operating_context = PUMP_OPERATING_CONTEXTS[
                (wo_index // len(PUMP_LOW_PRESSURE_CASES))
                % len(PUMP_OPERATING_CONTEXTS)
            ]
            description = (
                f"{pump_low_pressure_case['description']} on "
                f"{asset['description']} during {operating_context}"
            )
        age_days = random.randint(0, 900)
        reportdate = now - timedelta(days=age_days, hours=random.randint(0, 12))
        status = random.choices(statuses, [5, 8, 5, 7, 10, 8, 52, 5])[0]
        priority = random.choices([1, 2, 3, 4], [4, 18, 52, 26])[0]
        worktype = default_type if random.random() < 0.8 else random.choice(["CM", "PM", "INSP", "EM", "CAL"])
        lead = random.choice(tables["labor"])
        jp_candidates = [jp for jp in job_plans if jp["asset_class"] == asset["asset_class"]]
        jp = random.choice(jp_candidates) if jp_candidates and worktype in {"PM", "INSP", "CAL"} else None
        estimated = jp["estimated_hours"] if jp else round(random.uniform(1, 40), 1)
        target_start = reportdate + timedelta(days=max(0, 5 - priority) * random.randint(1, 4))
        actual_start = None
        actual_finish = None
        if status in {"INPRG", "COMP", "CLOSE"}:
            actual_start = target_start + timedelta(hours=random.randint(-8, 72))
        if status in completed_statuses:
            actual_finish = actual_start + timedelta(hours=max(1, int(estimated * random.uniform(0.6, 1.7))))
        wonum = f"WO{wo_index:07d}"
        failure_class = None if worktype in {"PM", "INSP"} else random.choice(["MECHANICAL", "ELECTRICAL", "CONTROL", "STRUCTURAL"])
        if pump_low_pressure_case and worktype not in {"PM", "INSP"}:
            failure_class = "MECHANICAL"
        long_description = (
            pump_low_pressure_case["details"]
            if pump_low_pressure_case
            else f"{description}. Verify as-found condition, complete required testing, and document as-left condition. Synthetic training record."
        )
        table_row = {
            "wonum": wonum,
            "siteid": asset["siteid"],
            "orgid": "OGE",
            "description": description,
            "long_description": long_description,
            "status": status,
            "worktype": worktype,
            "priority": priority,
            "assetnum": asset["assetnum"],
            "location": asset["location"],
            "reportedby": lead["laborcode"],
            "supervisor": random.choice(tables["labor"])["laborcode"],
            "lead": lead["laborcode"],
            "crewnum": lead["crewnum"],
            "reportdate": iso(reportdate),
            "target_start": iso(target_start),
            "target_finish": iso(target_start + timedelta(hours=estimated)),
            "actual_start": iso(actual_start),
            "actual_finish": iso(actual_finish),
            "estimated_labor_hours": estimated,
            "actual_labor_hours": round(estimated * random.uniform(0.5, 1.8), 1) if status in completed_statuses else None,
            "estimated_cost_usd": round(estimated * 110 + random.uniform(0, 5000), 2),
            "actual_cost_usd": round(estimated * random.uniform(90, 180) + random.uniform(0, 7000), 2) if status in completed_statuses else None,
            "jpnum": jp["jpnum"] if jp else None,
            "failure_class": failure_class,
            "problem_code": (
                pump_low_pressure_case["problem_code"]
                if pump_low_pressure_case and failure_class
                else random.choice(["LEAK", "VIBRATION", "OVERHEAT", "FAIL_TO_OPERATE", "DEGRADED_OUTPUT"]) if failure_class else None
            ),
            "cause_code": (
                pump_low_pressure_case["cause_code"]
                if pump_low_pressure_case and failure_class and status in completed_statuses
                else random.choice(["WEAR", "LOOSE_CONNECTION", "CONTAMINATION", "MISALIGNMENT", "AGE"]) if failure_class and status in completed_statuses else None
            ),
            "remedy_code": (
                pump_low_pressure_case["remedy_code"]
                if pump_low_pressure_case and failure_class and status in completed_statuses
                else random.choice(["ADJUSTED", "REPAIRED", "REPLACED", "CLEANED", "TESTED_OK"]) if failure_class and status in completed_statuses else None
            ),
            "parent": None,
            "synthetic": True,
        }
        tables["work_orders"].append(table_row)

        history_statuses = ["WAPPR"]
        if status != "WAPPR":
            history_statuses.append("APPR")
        if status in {"SCHED", "INPRG", "COMP", "CLOSE"}:
            history_statuses.append("SCHED")
        if status in {"INPRG", "COMP", "CLOSE"}:
            history_statuses.append("INPRG")
        if status in completed_statuses:
            history_statuses.extend(["COMP"] + (["CLOSE"] if status == "CLOSE" else []))
        elif status not in history_statuses:
            history_statuses.append(status)
        event_time = reportdate
        for sequence, history_status in enumerate(dict.fromkeys(history_statuses), 1):
            tables["work_order_status_history"].append(
                {
                    "wonum": wonum,
                    "siteid": asset["siteid"],
                    "sequence": sequence,
                    "status": history_status,
                    "changedate": iso(event_time),
                    "changeby": random.choice(tables["labor"])["laborcode"],
                    "memo": "Synthetic workflow status transition",
                }
            )
            event_time += timedelta(hours=random.randint(2, 72))

        if actual_start:
            assignments = random.sample(tables["labor"], k=random.randint(1, 3))
            for assignment_index, labor in enumerate(assignments, 1):
                regular_hours = round(random.uniform(1, max(2, estimated / len(assignments))), 1)
                tables["work_order_labor"].append(
                    {
                        "wonum": wonum,
                        "siteid": asset["siteid"],
                        "labor_sequence": assignment_index,
                        "laborcode": labor["laborcode"],
                        "craft": labor["craft"],
                        "startdate": iso(actual_start + timedelta(hours=assignment_index - 1)),
                        "regular_hours": regular_hours,
                        "premium_hours": round(regular_hours * 0.5, 1) if random.random() < 0.1 else 0.0,
                        "labor_rate_usd": round(random.uniform(58, 145), 2),
                    }
                )
        if status in completed_statuses and random.random() < 0.65:
            for material_index, item in enumerate(random.sample(tables["inventory_items"], k=random.randint(1, 4)), 1):
                quantity = random.randint(1, 6)
                tables["work_order_materials"].append(
                    {
                        "wonum": wonum,
                        "siteid": asset["siteid"],
                        "material_sequence": material_index,
                        "itemnum": item["itemnum"],
                        "quantity": quantity,
                        "unit": item["issueunit"],
                        "unitcost_usd": item["unitcost_usd"],
                        "linecost_usd": round(quantity * item["unitcost_usd"], 2),
                        "storeroom": f"{asset['siteid']}-MAIN",
                    }
                )

    for pm_index, asset in enumerate(random.sample(tables["assets"], 90), 1):
        jp_candidates = [jp for jp in job_plans if jp["asset_class"] == asset["asset_class"]]
        if not jp_candidates:
            continue
        frequency = random.choice([(1, "MONTHS"), (3, "MONTHS"), (6, "MONTHS"), (12, "MONTHS")])
        tables["preventive_maintenance"].append(
            {
                "pmnum": f"PM{pm_index:05d}",
                "description": f"{frequency[0]} {frequency[1].lower()} preventive maintenance - {asset['description']}",
                "siteid": asset["siteid"],
                "assetnum": asset["assetnum"],
                "location": asset["location"],
                "jpnum": random.choice(jp_candidates)["jpnum"],
                "frequency": frequency[0],
                "frequency_unit": frequency[1],
                "last_start_date": iso(now.date() - timedelta(days=random.randint(1, 365))),
                "next_due_date": iso(now.date() + timedelta(days=random.randint(5, 180))),
                "status": "ACTIVE",
                "use_target_date": True,
            }
        )

    meter_types = {
        "GASTURB": ("OPERATING_HOURS", "HOURS", 15_000, 85_000),
        "GENSET": ("RUN_HOURS", "HOURS", 100, 8_000),
        "BREAKER": ("OPERATIONS", "COUNT", 50, 5_000),
        "PUMP": ("OPERATING_HOURS", "HOURS", 5_000, 90_000),
        "TRANSFORMER": ("LOAD_MWH", "MWH", 50_000, 9_000_000),
    }
    for asset in tables["assets"]:
        meter_name, unit, low, high = meter_types.get(asset["asset_class"], ("OPERATING_HOURS", "HOURS", 500, 40_000))
        base = random.uniform(low, high)
        for reading_index in range(4):
            tables["meter_readings"].append(
                {
                    "assetnum": asset["assetnum"],
                    "siteid": asset["siteid"],
                    "metername": meter_name,
                    "readingdate": iso(now - timedelta(days=(3 - reading_index) * 30)),
                    "reading": round(base + reading_index * random.uniform(10, 500), 2),
                    "unit": unit,
                    "inspector": random.choice(tables["labor"])["laborcode"],
                }
            )

    document_rows = []
    for document in DOCUMENTS:
        document_rows.append(
            {
                "docnum": document["docnum"],
                "title": document["title"],
                "library": document["library"],
                "revision": document["revision"],
                "owner": document["owner"],
                "asset_class": document["asset_class"],
                "effective_date": "2026-01-15",
                "review_date": "2028-01-15",
                "status": "ACTIVE",
                "pdf_file": f"{document['docnum']}-{slug(document['title'])}.pdf",
                "docx_file": f"{document['docnum']}-{slug(document['title'])}.docx",
                "synthetic": True,
            }
        )
    tables["document_register"] = document_rows
    return tables


def add_docx_header(document: Document, item: dict[str, Any]) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("OGE OPERATIONS & MAINTENANCE")
    run.bold = True
    run.font.size = Pt(10)
    heading = document.add_heading(item["title"], level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata = document.add_table(rows=4, cols=4)
    metadata.style = "Table Grid"
    values = [
        ("Document", item["docnum"], "Revision", item["revision"]),
        ("Owner", item["owner"], "Status", "ACTIVE"),
        ("Effective", "January 15, 2026", "Review", "January 15, 2028"),
        ("Classification", "SYNTHETIC DEMO", "Asset class", item["asset_class"]),
    ]
    for row, values_row in zip(metadata.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = value
    document.add_paragraph()
    warning = document.add_paragraph()
    warning.add_run("TRAINING AND DEMONSTRATION USE ONLY. ").bold = True
    warning.add_run("This synthetic document does not replace approved company procedures or manufacturer instructions.")


def render_docx(item: dict[str, Any], output_path: Path) -> None:
    document = Document()
    add_docx_header(document, item)
    document.add_heading("Purpose", level=1)
    document.add_paragraph(item["purpose"])
    for title, bullets in item["sections"]:
        document.add_heading(title, level=1)
        for bullet in bullets:
            document.add_paragraph(bullet, style="List Bullet")
    document.add_heading("Records", level=1)
    document.add_paragraph("Record required readings, test results, deficiencies, labor, materials, and closeout notes in the associated work order.")
    document.add_heading("Revision history", level=1)
    table = document.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("Revision", "Date", "Author", "Description")):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, (item["revision"], "2026-01-15", item["owner"], "Synthetic demo issue")):
        cell.text = value
    document.save(output_path)


def render_pdf(item: dict[str, Any], output_path: Path) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleCenter", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#17365D"))
    heading_style = ParagraphStyle("Section", parent=styles["Heading2"], textColor=colors.HexColor("#1F4E79"), spaceBefore=12)
    body = [Paragraph("OGE OPERATIONS &amp; MAINTENANCE", styles["Heading4"]), Paragraph(item["title"], title_style), Spacer(1, 0.15 * inch)]
    metadata = [
        ["Document", item["docnum"], "Revision", item["revision"]],
        ["Owner", item["owner"], "Status", "ACTIVE"],
        ["Effective", "2026-01-15", "Review", "2028-01-15"],
        ["Classification", "SYNTHETIC DEMO", "Asset class", item["asset_class"]],
    ]
    table = Table(metadata, colWidths=[0.9 * inch, 2.3 * inch, 0.9 * inch, 2.3 * inch])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#D9EAF7")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#D9EAF7")), ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    body.extend([table, Spacer(1, 0.15 * inch), Paragraph("<b>TRAINING AND DEMONSTRATION USE ONLY.</b> This synthetic document does not replace approved company procedures or manufacturer instructions.", styles["BodyText"]), Paragraph("Purpose", heading_style), Paragraph(item["purpose"], styles["BodyText"])])
    for title, bullets in item["sections"]:
        body.append(Paragraph(title, heading_style))
        for bullet in bullets:
            body.append(Paragraph(f"&#8226; {bullet}", styles["BodyText"]))
            body.append(Spacer(1, 0.04 * inch))
    body.extend([Paragraph("Records", heading_style), Paragraph("Record required readings, test results, deficiencies, labor, materials, and closeout notes in the associated work order.", styles["BodyText"]), PageBreak(), Paragraph("Revision history", heading_style)])
    history = Table([["Revision", "Date", "Author", "Description"], [item["revision"], "2026-01-15", item["owner"], "Synthetic demo issue"]], colWidths=[0.8 * inch, 1.2 * inch, 2 * inch, 2.5 * inch])
    history.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    body.append(history)
    pdf = SimpleDocTemplate(str(output_path), pagesize=letter, rightMargin=0.6 * inch, leftMargin=0.6 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=item["title"], author="OGE Synthetic Demo")
    pdf.build(body)


def generate_documents() -> None:
    for item in DOCUMENTS:
        directory = SOP_DIR if item["library"] == "Standard Operating Procedures" else MANUAL_DIR
        base_name = f"{item['docnum']}-{slug(item['title'])}"
        render_docx(item, directory / f"{base_name}.docx")
        render_pdf(item, directory / f"{base_name}.pdf")


def main() -> None:
    for directory in (PARQUET_DIR, SOP_DIR, MANUAL_DIR):
        directory.mkdir(parents=True, exist_ok=True)
    tables = generate_data()
    for name, rows in tables.items():
        write_table(name, rows)
    generate_documents()
    manifest = {
        "generated_at": datetime.now(UTC).isoformat(),
        "seed": SEED,
        "classification": "SYNTHETIC",
        "tables": {name: len(rows) for name, rows in tables.items()},
        "documents": len(DOCUMENTS),
        "document_files": len(DOCUMENTS) * 2,
    }
    (OUTPUT_ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
